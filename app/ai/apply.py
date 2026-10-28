"""Apply accepted proposals to a document, producing the bytes of the next version."""

from dataclasses import dataclass
from pathlib import Path

import docx
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from app.ai.evidence import normalize
from app.ai.render_docx import _save, render_content
from app.ai.types import ExtractedContent


@dataclass
class Edit:
    current: str
    proposed: str


@dataclass
class ApplyResult:
    data: bytes
    applied: int
    appended: int


def _set_text(paragraph: Paragraph, text: str) -> None:
    """Replace text but keep the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _replace(paragraph: Paragraph, edit: Edit) -> bool:
    text = paragraph.text
    if not text.strip():
        return False
    if edit.current in text:
        _set_text(paragraph, text.replace(edit.current, edit.proposed, 1))
        return True
    if normalize(edit.current) and normalize(edit.current) == normalize(text):
        _set_text(paragraph, edit.proposed)
        return True
    return False


def _replace_row(cells: list[_Cell], edit: Edit) -> bool:
    """Table rows were extracted as 'cell | cell'; edits spanning a row are matched cell by cell."""
    unique = [c for i, c in enumerate(cells) if c not in cells[:i]]
    joined = " | ".join(c.text.strip() for c in unique if c.text.strip())
    if normalize(joined) != normalize(edit.current):
        return False
    filled = [c for c in unique if c.text.strip()]
    parts = [p.strip() for p in edit.proposed.split(" | ")]
    if len(parts) != len(filled):
        return False  # row shape changed; the edit goes to the Revisions appendix instead of guessing
    for cell, value in zip(filled, parts, strict=True):
        if cell.text.strip() != value:
            _set_text(cell.paragraphs[0], value)
            for extra in cell.paragraphs[1:]:
                _set_text(extra, "")
    return True


def apply_docx(path: Path, edits: list[Edit]) -> ApplyResult:
    document = docx.Document(str(path))
    paragraphs = list(document.paragraphs)
    rows = []
    for table in document.tables:
        for row in table.rows:
            rows.append(list(row.cells))
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    applied, leftovers = 0, []
    for edit in edits:
        done = any(_replace_row(cells, edit) for cells in rows) if " | " in edit.current else False
        done = done or any(_replace(p, edit) for p in paragraphs)
        if done:
            applied += 1
        else:
            leftovers.append(edit.proposed)
    if leftovers:
        document.add_heading("Revisions", level=1)
        for text in leftovers:
            document.add_paragraph(text)
    return ApplyResult(_save(document), applied, len(leftovers))


def apply_content(content: ExtractedContent, edits: list[Edit]) -> ApplyResult:
    """For sources that cannot be edited in place (PDF): rewrite paragraphs and render a new .docx."""
    paragraphs = list(content.paragraphs)
    applied, leftovers = 0, []
    for edit in edits:
        for i, p in enumerate(paragraphs):
            if edit.current in p or normalize(edit.current) == normalize(p):
                paragraphs[i] = (
                    p.replace(edit.current, edit.proposed, 1) if edit.current in p else edit.proposed
                )
                applied += 1
                break
        else:
            leftovers.append(edit.proposed)
    updated = content.model_copy(update={"paragraphs": paragraphs})
    return ApplyResult(render_content(updated, appendix=leftovers), applied, len(leftovers))
