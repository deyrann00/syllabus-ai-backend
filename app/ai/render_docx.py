"""Generate .docx files: wizard drafts in university format, and re-rendered PDFs after revisions."""

import io

import docx
from docx.document import Document as DocxDocument

from app.ai.types import ExtractedContent

# Wizard field → (heading, label). Headings use the section names the extractor recognises.
WIZARD_LAYOUT: list[tuple[str, list[tuple[str, str]]]] = [
    (
        "Course information",
        [
            ("title", "Course title"),
            ("credits", "Credits"),
            ("language", "Language"),
            ("semester", "Semester"),
            ("degree_level", "Degree level"),
        ],
    ),
    (
        "Academic context",
        [
            ("program", "Educational program"),
            ("programRole", "Role in the program"),
            ("prerequisites", "Prerequisites"),
        ],
    ),
    ("Course purpose", [("purpose", "")]),
    ("Learning outcomes", [("outcomes", "")]),
    ("Content", [("modules", "Modules"), ("weeks", "")]),
    ("Teaching methods", [("methods", "")]),
    ("Assessment", [("assignments", ""), ("final", "")]),
    ("Resources", [("resources", "")]),
    (
        "Course policy",
        [("attendance", "Attendance"), ("integrity", "Academic integrity"), ("appeals", "Appeals")],
    ),
]


def _save(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_wizard(fields: dict[str, str], *, course_name: str) -> bytes:
    document = docx.Document()
    document.add_heading(f"Syllabus: {fields.get('title') or course_name}", level=0)
    for heading, items in WIZARD_LAYOUT:
        document.add_heading(heading, level=1)
        wrote = False
        for key, label in items:
            value = (fields.get(key) or "").strip()
            if not value:
                continue
            wrote = True
            lines = [ln.strip() for ln in value.splitlines() if ln.strip()]
            if label and len(lines) == 1:
                document.add_paragraph(f"{label}: {lines[0]}")
                continue
            if label:
                document.add_paragraph(f"{label}:")
            for line in lines:
                document.add_paragraph(line)
        if not wrote:
            document.add_paragraph("To be completed.")
    return _save(document)


def render_content(content: ExtractedContent, *, appendix: list[str] | None = None) -> bytes:
    """Re-create a document from extracted paragraphs (used when the source was a PDF)."""
    document = docx.Document()
    headings = set(content.headings)
    for paragraph in content.paragraphs:
        if paragraph in headings:
            document.add_heading(paragraph, level=1)
        else:
            document.add_paragraph(paragraph)
    if appendix:
        document.add_heading("Revisions", level=1)
        for text in appendix:
            document.add_paragraph(text)
    return _save(document)
