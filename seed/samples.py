"""Synthetic syllabus documents for the demo seed and tests.

`with_issues=True` reproduces the prototype's sample review: LO2 uses "understand" (3.5), Week 9 maps to no
outcome (1.2), deployment is never taught (4.3), LO5 is not assessed (6.1), the rubric has no performance
levels (6.4) and group work lacks individual responsibility (9.2).
"""

import io

import docx


def sample_syllabus(course: str, *, with_issues: bool, language_note: str = "English") -> bytes:
    d = docx.Document()
    d.add_heading(f"Syllabus: {course}", level=0)

    d.add_heading("Course information", level=1)
    d.add_paragraph(f"Course title: {course}. Credits: 5. Language of instruction: {language_note}.")
    d.add_heading("Course purpose", level=1)
    d.add_paragraph(f"The course introduces the core methods of {course.lower()} and their responsible use.")

    d.add_heading("Learning outcomes", level=1)
    d.add_paragraph("LO1. Explain the main families of methods and select one for a given problem.")
    if with_issues:
        d.add_paragraph("LO2. Students will understand the fundamental principles of machine learning.")
    else:
        d.add_paragraph(
            "LO2. Compare machine-learning methods and justify their selection for a given problem."
        )
    d.add_paragraph("LO3. Implement and train a model on a real dataset.")
    d.add_paragraph("LO4. Deploy a trained model as a simple service.")
    d.add_paragraph("LO5. Evaluate ethical risks of an AI solution.")

    d.add_heading("Content", level=1)
    topics = [
        "Introduction and problem framing (LO1)",
        "Search and optimisation (LO1)",
        "Probabilistic reasoning (LO1)",
        "Supervised learning (LO2)",
        "Model selection (LO2)",
        "Unsupervised learning (LO2)",
        "Feature engineering lab (LO3)",
        "Training at scale (LO3)",
        "Neural networks lab" if with_issues else "Neural networks lab (LO3)",
        "Responsible AI and ethics (LO5)",
        "Project work (LO3)",
        "Model evaluation" if with_issues else "Model evaluation and deployment (LO4)",
    ]
    for week, topic in enumerate(topics, start=1):
        d.add_paragraph(f"Week {week}: {topic}.")

    d.add_heading("Teaching methods", level=1)
    d.add_paragraph(
        "Lectures, weekly labs, guided project work and peer review sessions support active learning."
    )

    d.add_heading("Assessment", level=1)
    table = d.add_table(rows=0, cols=3)
    rows = [("Midterm exam", "30%", "LO1, LO2"), ("Lab portfolio", "30%", "LO3")]
    rows.append(
        ("Final project", "40%", "LO3, LO4") if with_issues else ("Final project", "40%", "LO3, LO4, LO5")
    )
    for cells in rows:
        row = table.add_row().cells
        for cell, value in zip(row, cells, strict=True):
            cell.text = value
    if with_issues:
        d.add_paragraph("Project assessed on technical quality and presentation.")
    else:
        d.add_paragraph(
            "Project rubric: technical quality 60%, evaluation 25%, presentation 15%, with level descriptors for "
            "excellent, satisfactory and insufficient work."
        )

    d.add_heading("Resources", level=1)
    d.add_paragraph("Russell S., Norvig P. Artificial Intelligence: A Modern Approach. 4th ed., 2021.")
    d.add_paragraph("Géron A. Hands-On Machine Learning. 3rd ed., 2022.")

    d.add_heading("Course policy", level=1)
    d.add_paragraph("Attendance: students attend at least 80% of classes; absences are agreed in advance.")
    d.add_paragraph("Academic integrity: plagiarism and unauthorised AI use lead to a zero for the task.")
    if with_issues:
        d.add_paragraph("Students are expected to communicate respectfully.")
    else:
        d.add_paragraph(
            "In group work, each student documents their individual contribution and responsibility, "
            "and communicates respectfully with peers."
        )
    d.add_paragraph("Appeals: students may request a grade review within five working days.")

    buffer = io.BytesIO()
    d.save(buffer)
    return buffer.getvalue()


BROKEN_PDF = b"%PDF-1.4\n% truncated upload\n"
